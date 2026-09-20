"""
modules/excel_export.py
Generate styled Excel reports using openpyxl.
Supports attendance, eligibility, and marks sheets.
"""
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import (PatternFill, Font, Alignment,
                              Border, Side, GradientFill)
from openpyxl.utils  import get_column_letter
import config

# ── Colour palette ────────────────────────────────────────────────────────────
NAVY   = "1A237E"
WHITE  = "FFFFFF"
GREEN  = "C8E6C9"
RED    = "FFCDD2"
AMBER  = "FFF9C4"
LGREY  = "F5F5F5"
DKGREY = "424242"


def _header_style(ws, row: int, cols: list[str], fill_hex: str = NAVY):
    fill = PatternFill("solid", fgColor=fill_hex)
    font = Font(bold=True, color=WHITE, size=11)
    for col, text in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _thin_border():
    s = Side(style="thin", color="BDBDBD")
    return Border(left=s, right=s, top=s, bottom=s)


def _auto_width(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 4


def _title_row(ws, title: str, ncols: int):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = Font(bold=True, size=14, color=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    sub = ws.cell(row=2, column=1,
                  value=f"Generated: {datetime.now().strftime('%d %b %Y  %H:%M')}")
    sub.font = Font(italic=True, size=10, color=DKGREY)
    sub.alignment = Alignment(horizontal="center")


# ─── Attendance Export ────────────────────────────────────────────────────────

def export_attendance(records: list[dict], subject: str,
                      date_range: str = "") -> str:
    """
    Export attendance records to a styled Excel file.
    Returns absolute path to the created file.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"

    title = f"Attendance Report — {subject}" + (f"  ({date_range})" if date_range else "")
    cols  = ["#", "USN", "Student Name", "Subject", "Date", "Time", "Status", "Confidence"]
    _title_row(ws, title, len(cols))
    _header_style(ws, 3, cols)

    for i, r in enumerate(records, start=1):
        row_n = i + 3
        status = r.get("status", "")
        fill   = PatternFill("solid", fgColor=GREEN if status == "Present" else RED)
        values = [
            i,
            r.get("usn", ""),
            r.get("name", ""),
            r.get("subject", ""),
            r.get("date", ""),
            str(r.get("timestamp", ""))[:8],
            status,
            f"{round(r.get('confidence', 0) * 100, 1)}%",
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=row_n, column=col, value=val)
            cell.fill = fill
            cell.border = _thin_border()
            cell.alignment = Alignment(horizontal="center")

    _auto_width(ws)
    ws.freeze_panes = "A4"

    fname = f"attendance_{subject.replace(' ','_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path  = os.path.join(config.EXPORTS_DIR, fname)
    wb.save(path)
    return path


# ─── Eligibility Export ───────────────────────────────────────────────────────

def export_eligibility(report: dict) -> str:
    """
    Export eligibility report for one subject.
    Returns path to saved file.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Eligibility"

    subject   = report.get("subject", "")
    threshold = report.get("threshold", 75)
    cols = ["#", "USN", "Student Name", "Present", "Total", "Attendance %", "Status"]
    _title_row(ws, f"Internal Eligibility — {subject}  (Threshold: {threshold}%)", len(cols))
    _header_style(ws, 3, cols)

    all_rows = (
        [(r, "Eligible")   for r in report.get("eligible",   [])] +
        [(r, "At Risk")    for r in report.get("at_risk",    [])] +
        [(r, "Ineligible") for r in report.get("ineligible", [])]
    )

    color_map = {"Eligible": GREEN, "At Risk": AMBER, "Ineligible": RED}

    for i, (r, status) in enumerate(all_rows, start=1):
        row_n = i + 3
        fill  = PatternFill("solid", fgColor=color_map[status])
        values = [i, r.get("usn",""), r.get("name",""),
                  r.get("present",0), r.get("total",0),
                  f"{r.get('pct',0):.1f}%", status]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=row_n, column=col, value=val)
            cell.fill   = fill
            cell.border = _thin_border()
            cell.alignment = Alignment(horizontal="center")

    _auto_width(ws)
    fname = f"eligibility_{subject.replace(' ','_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path  = os.path.join(config.EXPORTS_DIR, fname)
    wb.save(path)
    return path


# ─── Marks Export ─────────────────────────────────────────────────────────────

def export_marks(records: list[dict], subject: str = "All") -> str:
    """Export IA marks to Excel with totals and averages."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Marks"

    cols = ["#", "USN", "Student Name", "Subject", "IA1", "IA2", "IA3", "Total", "Avg"]
    _title_row(ws, f"Internal Assessment Marks — {subject}", len(cols))
    _header_style(ws, 3, cols)

    row_fill = PatternFill("solid", fgColor=LGREY)
    for i, r in enumerate(records, start=1):
        row_n  = i + 3
        values = [
            i, r.get("usn",""), r.get("name",""), r.get("subject",""),
            r.get("ia1", "—"), r.get("ia2", "—"), r.get("ia3", "—"),
            r.get("total", 0), r.get("avg", 0),
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=row_n, column=col, value=val)
            cell.border = _thin_border()
            cell.alignment = Alignment(horizontal="center")
            if i % 2 == 0:
                cell.fill = row_fill

    _auto_width(ws)
    fname = f"marks_{subject.replace(' ','_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path  = os.path.join(config.EXPORTS_DIR, fname)
    wb.save(path)
    return path
